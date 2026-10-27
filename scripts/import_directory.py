"""Turn a filled import template (docs/import/abyad-import-template.xlsx) into one SQL script.

  python3 -I scripts/import_directory.py filled.xlsx --admin-email owner@example.com [--out import.sql] [--no-resolve]

1. Reads the three sheets, checks every row (required cells, dropdown values, emails, prices,
   hours, that every code/email a row points to exists) and prints an Arabic report.
2. Follows short Google Maps links to their coordinates (the same rules as the app's map-link).
3. Only when the whole file is clean, writes the SQL. The SQL runs in ONE transaction as the
   given platform admin through the platform's own admin functions (same checks and audit as
   the admin screens). It is safe to run twice: existing complexes, branches, places, accounts
   and memberships are reused, prices and weekly hours are replaced.

Accounts are created without a password: each person signs in with Google on the same email,
or uses «نسيت كلمة المرور» to choose one (the Arabic reset email).
"""
import argparse
import json
import re
import sys
import urllib.parse
import urllib.request
from openpyxl import load_workbook

TYPES = {'عيادة': 'clinic', 'صيدلية': 'pharmacy', 'مختبر': 'laboratory'}
ROLES = {'سكرتير': 'receptionist', 'صيدلاني': 'pharmacist', 'موظف مختبر': 'lab_technician', 'معتمد نتائج المختبر': 'lab_approver'}
ROLE_FOR_TYPE = {'receptionist': 'clinic', 'pharmacist': 'pharmacy', 'lab_technician': 'laboratory', 'lab_approver': 'laboratory'}
DAYS = {'السبت': 6, 'الأحد': 0, 'الإثنين': 1, 'الثلاثاء': 2, 'الأربعاء': 3, 'الخميس': 4, 'الجمعة': 5}  # 0 = Sunday (schedules.weekday)
EMAIL = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
SPAN = re.compile(r'^([01]?\d|2[0-3]):([0-5]\d)\s*-\s*([01]?\d|2[0-3]):([0-5]\d)$')
NUM = r'(-?\d{1,3}(?:\.\d+)?)'
MAP_PATTERNS = [rf'!3d{NUM}!4d{NUM}', rf'[?&](?:q|query|ll|sll|daddr|destination|center|mlat)=(?:loc:)?{NUM}\s*,\s*{NUM}',
                rf'@{NUM},{NUM}', rf'^geo:{NUM},{NUM}', rf'^{NUM}\s*[,،\s]\s*{NUM}$']
EXAMPLE = re.compile(r'example\.com|xxxxxxxx|yyyyyyyy', re.I)

errors: list[str] = []


def err(sheet, row, msg):
    errors.append(f'• {sheet} — سطر {row}: {msg}')


def cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def rows(wb, name, ncols):
    if name not in wb.sheetnames:
        errors.append(f'• الورقة «{name}» غير موجودة بالملف')
        return []
    out = []
    for r, values in enumerate(wb[name].iter_rows(min_row=3, max_col=ncols, values_only=True), start=3):
        vals = [cell(v) for v in values]
        if any(vals):
            if any(EXAMPLE.search(v) for v in vals):
                err(name, r, 'هذا سطر مثال — امسحه أو اكتب بياناتك فوقه')
                continue
            out.append((r, vals))
    return out


def parse_map(s):
    try:
        s = urllib.parse.unquote(s.strip())
    except Exception:
        pass
    for p in MAP_PATTERNS:
        m = re.search(p, s)
        if m:
            lat, lng = float(m.group(1)), float(m.group(2))
            if abs(lat) <= 90 and abs(lng) <= 180 and not (lat == 0 and lng == 0):
                return round(lat, 6), round(lng, 6)
    lat, lon = re.search(rf'[?&]mlat={NUM}', s), re.search(rf'[?&]mlon={NUM}', s)
    return (round(float(lat.group(1)), 6), round(float(lon.group(1)), 6)) if lat and lon else None


def resolve_map(link, follow):
    p = parse_map(link)
    if p or not follow or not re.match(r'^https?://', link):
        return p
    url = link
    for _ in range(6):  # follow redirects one hop at a time, checking each URL
        req = urllib.request.Request(url, method='GET', headers={'User-Agent': 'Mozilla/5.0'})
        opener = urllib.request.build_opener(NoRedirect)
        try:
            resp = opener.open(req, timeout=15)
            body = resp.read(200_000).decode('utf-8', 'ignore')
            return parse_map(body) if not parse_map(url) else parse_map(url)
        except urllib.error.HTTPError as e:
            loc = e.headers.get('Location')
            if not loc:
                return None
            url = urllib.parse.urljoin(url, loc)
            p = parse_map(url)
            if p:
                return p
        except Exception:
            return None
    return None


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


def q(v):
    """SQL literal (NULL for empty)."""
    if v is None or v == '':
        return 'null'
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--admin-email', required=True)
    ap.add_argument('--out', default='import.sql')
    ap.add_argument('--no-resolve', action='store_true', help='do not follow short map links (offline check)')
    a = ap.parse_args()
    wb = load_workbook(a.xlsx, data_only=True)

    # ---------------- places
    places = {}
    for r, v in rows(wb, 'الجهات', 11):
        code, name, typ, complex_, branch, city, area, address, link, phone, desc = v
        if not code: err('الجهات', r, 'رمز الجهة فارغ'); continue
        if code in places: err('الجهات', r, f'الرمز {code} مكرر'); continue
        if not name: err('الجهات', r, 'اسم الجهة فارغ')
        if typ not in TYPES: err('الجهات', r, 'النوع لازم يكون: عيادة أو صيدلية أو مختبر')
        if not city: err('الجهات', r, 'المحافظة فارغة')
        if not area: err('الجهات', r, 'المنطقة فارغة')
        loc = None
        if not link:
            err('الجهات', r, 'رابط الموقع فارغ — بدونه ما يطلع للمرضى «الأقرب إليك»')
        else:
            loc = resolve_map(link, not a.no_resolve)
            if not loc and not a.no_resolve:
                err('الجهات', r, 'ما قدرنا نقرا الموقع من الرابط — انسخه من «مشاركة» بخرائط Google')
        places[code] = dict(row=r, name=name, type=TYPES.get(typ), complex=complex_, branch=branch or 'الفرع الرئيسي',
                            city=city, area=area, address=address, loc=loc, phone=phone, desc=desc)

    # ---------------- doctors
    doctors, emails = [], {}
    for r, v in rows(wb, 'الأطباء', 19):
        code, name, title, spec, email, phone, lic, bio, p1, d1, p2, d2, *days = v
        email = email.lower()
        if code not in places: err('الأطباء', r, f'رمز الجهة «{code}» مو موجود بورقة الجهات'); continue
        if places[code]['type'] != 'clinic': err('الأطباء', r, f'الجهة {code} مو عيادة')
        if not name: err('الأطباء', r, 'اسم الطبيب فارغ')
        if not spec: err('الأطباء', r, 'الاختصاص فارغ')
        if not EMAIL.match(email): err('الأطباء', r, 'البريد الإلكتروني غير صحيح')
        elif emails.get(email, ('doctor', name))[1] != name: err('الأطباء', r, f'نفس البريد {email} مكتوب لشخص ثاني')
        for label, price, mins, needed in (('الكشفية', p1, d1, True), ('المراجعة', p2, d2, False)):
            if (price or mins) or needed:
                if not price.isdigit(): err('الأطباء', r, f'سعر {label} لازم رقم بالدينار، مثل 25000')
                if not mins.isdigit() or not 5 <= int(mins or 0) <= 240: err('الأطباء', r, f'مدة {label} لازم دقائق بين 5 و 240')
        slots = []
        for day, val in zip(DAYS, days):
            for part in [x for x in re.split(r'[,،]', val) if x.strip()]:
                m = SPAN.match(part.strip())
                if not m: err('الأطباء', r, f'دوام {day} «{part.strip()}» لازم يكون مثل 16:00-21:00'); continue
                s, e = f'{int(m[1]):02d}:{m[2]}', f'{int(m[3]):02d}:{m[4]}'
                if s >= e: err('الأطباء', r, f'دوام {day}: وقت النهاية قبل البداية')
                slots.append({'weekday': DAYS[day], 'start': s, 'end': e})
        if not slots: err('الأطباء', r, 'ما مكتوب أي دوام — المرضى ما يكدرون يحجزون')
        emails.setdefault(email, ('doctor', name))
        doctors.append(dict(row=r, code=code, name=name, title=title, spec=spec, email=email, lic=lic, bio=bio,
                            services=[s for s in (('كشفية', p1, d1), ('مراجعة', p2, d2)) if s[1].isdigit() and s[2].isdigit()], slots=slots))

    # ---------------- staff
    staff = []
    doctor_emails = {d['email'] for d in doctors}
    for r, v in rows(wb, 'الموظفون', 6):
        code, name, role, email, phone, doc = v
        email, doc = email.lower(), doc.lower()
        if code not in places: err('الموظفون', r, f'رمز الجهة «{code}» مو موجود بورقة الجهات'); continue
        if role not in ROLES: err('الموظفون', r, 'الدور لازم من القائمة'); continue
        if places[code]['type'] != ROLE_FOR_TYPE[ROLES[role]]: err('الموظفون', r, f'دور «{role}» ما يناسب نوع الجهة {code}')
        if not name: err('الموظفون', r, 'الاسم فارغ')
        if not EMAIL.match(email): err('الموظفون', r, 'البريد الإلكتروني غير صحيح')
        elif email in emails and emails[email][1] != name: err('الموظفون', r, f'نفس البريد {email} مكتوب لشخص ثاني')
        if ROLES[role] == 'receptionist':
            if not doc: err('الموظفون', r, 'السكرتير يحتاج بريد الطبيب اللي يشتغل وياه')
            elif doc not in doctor_emails: err('الموظفون', r, f'بريد الطبيب {doc} مو موجود بورقة الأطباء')
        emails.setdefault(email, ('staff', name))
        staff.append(dict(row=r, code=code, name=name, role=ROLES[role], email=email, doctor=doc or None))

    print(f'الجهات: {len(places)} · الأطباء: {len(doctors)} · الموظفون: {len(staff)} · الحسابات: {len(emails)}')
    if errors:
        print(f'\nلازم نصلح {len(errors)} شي قبل الإدخال:')
        print('\n'.join(errors))
        sys.exit(1)
    if not places:
        print('الملف فارغ.'); sys.exit(1)

    # ---------------- SQL
    S = ['-- Generated by scripts/import_directory.py — one transaction, through the admin functions.', 'begin;',
         'create temp table imp_place (code text primary key, facility_id uuid, organization_id uuid) on commit drop;',
         'create temp table imp_doc (email text primary key, practitioner_id uuid) on commit drop;',
         f"""select set_config('request.jwt.claims', json_build_object('sub', (select u.id from auth.users u join public.platform_admins pa on pa.user_id = u.id and pa.revoked_at is null where lower(u.email) = lower({q(a.admin_email)})), 'role', 'authenticated', 'aal', 'aal2', 'email', {q(a.admin_email)})::text, true);""",
         "do $imp$ begin if auth.uid() is null then raise exception 'the admin email is not an active platform admin'; end if; end $imp$;"]
    # accounts (no password)
    for email, (_, name) in emails.items():
        S.append(f"""insert into auth.users (instance_id, id, aud, role, email, encrypted_password, email_confirmed_at, raw_app_meta_data, raw_user_meta_data, created_at, updated_at, confirmation_token, recovery_token, email_change_token_new, email_change)
select '00000000-0000-0000-0000-000000000000', gen_random_uuid(), 'authenticated', 'authenticated', {q(email)}, '', now(), '{{"provider":"email","providers":["email"]}}', jsonb_build_object('display_name', {q(name)}), now(), now(), '', '', '', ''
 where not exists (select 1 from auth.users where lower(email) = {q(email)});
insert into auth.identities (id, user_id, provider_id, identity_data, provider, created_at, updated_at)
select gen_random_uuid(), u.id, u.id::text, jsonb_build_object('sub', u.id::text, 'email', u.email), 'email', now(), now()
  from auth.users u where lower(u.email) = {q(email)} and not exists (select 1 from auth.identities i where i.user_id = u.id and i.provider = 'email');""")
    for code, p in places.items():
        kind = 'medical_complex' if p['complex'] else 'independent'
        org_name = p['complex'] or p['name']
        lat, lng = p['loc'] or (None, None)
        S.append(f"""do $imp$ declare o uuid; b uuid; f uuid; c uuid; begin
  select id into c from public.cities where name_ar = {q(p['city'])};
  select id into o from public.organizations where name = {q(org_name)} and kind::text = {q(kind)};
  if o is null then o := (public.admin_create_organization({q(org_name)}, {q(kind)}, true)).id; end if;
  select id into b from public.branches where organization_id = o and name = {q(p['branch'])};
  if b is null then b := (public.admin_create_branch(o, {q(p['branch'])}, c, {q(p['area'])}, {q(p['address'])}, {q(lat)}, {q(lng)})).id; end if;
  select id into f from public.facilities where branch_id = b and name = {q(p['name'])};
  if f is null then f := (public.admin_create_facility(o, b, {q(p['type'])}::public.facility_type, {q(p['name'])}, {q(p['desc'])}, {q(p['phone'])}, true)).id; end if;
  insert into pg_temp.imp_place values ({q(code)}, f, o);
end $imp$;""")
    for d in doctors:
        S.append(f"""do $imp$ declare pr uuid; uid uuid; pl pg_temp.imp_place; sp uuid; admin_claims text; begin
  select * into pl from pg_temp.imp_place where code = {q(d['code'])};
  select id into uid from auth.users where lower(email) = {q(d['email'])};
  select id into sp from public.specialties where name_ar = {q(d['spec'])};
  pr := (public.admin_upsert_practitioner_by_email({q(d['email'])}, {q(d['name'])}, sp, {q(d['title'])}, {q(d['lic'])}, {q(d['bio'])}, false)).id;
  insert into pg_temp.imp_doc values ({q(d['email'])}, pr) on conflict do nothing;
  if not exists (select 1 from public.memberships m where m.user_id = uid and m.facility_id = pl.facility_id and m.role = 'doctor' and m.status <> 'revoked') then
    perform public.admin_grant_membership(uid, pl.organization_id, pl.facility_id, 'doctor');
  end if;
  -- prices and hours are set the way the doctor sets them (the doctor's own schedule rights);
  -- the import itself is audited under the admin at the end
  admin_claims := current_setting('request.jwt.claims', true);
  perform set_config('request.jwt.claims', json_build_object('sub', uid, 'role', 'authenticated', 'aal', 'aal1')::text, true);""" + ''.join(f"""
  perform public.upsert_clinic_service(pl.facility_id, {q(n)}, {int(m)}, {int(p)}, pr, 0, null, null, null,
    (select s.id from public.clinic_services s where s.facility_id = pl.facility_id and s.practitioner_id = pr and s.name = {q(n)} limit 1));"""
                    for n, p, m in d['services']) + f"""
  perform public.set_weekly_schedule(pl.facility_id, pr, {q(json.dumps(d['slots']))}::jsonb);
  perform set_config('request.jwt.claims', admin_claims, true);
end $imp$;""")
    for s in staff:
        S.append(f"""do $imp$ declare uid uuid; pl pg_temp.imp_place; mem uuid; begin
  select * into pl from pg_temp.imp_place where code = {q(s['code'])};
  select id into uid from auth.users where lower(email) = {q(s['email'])};
  select id into mem from public.memberships m where m.user_id = uid and m.facility_id = pl.facility_id and m.role = {q(s['role'])}::public.member_role and m.status <> 'revoked';
  if mem is null then mem := (public.admin_grant_membership(uid, pl.organization_id, pl.facility_id, {q(s['role'])}::public.member_role)).id; end if;""" + (f"""
  perform public.set_reception_scope(mem, (select practitioner_id from pg_temp.imp_doc where email = {q(s['doctor'])}), true);""" if s['doctor'] else '') + """
end $imp$;""")
    S.append(f"""select private.audit('directory.imported', 'platform', null, null, null, {q(a.xlsx.split('/')[-1])},
  jsonb_build_object('places', {len(places)}, 'doctors', {len(doctors)}, 'staff', {len(staff)}, 'accounts', {len(emails)}));""")
    S.append("select (select count(*) from pg_temp.imp_place) as places, (select count(*) from pg_temp.imp_doc) as doctors;")
    S.append('commit;')
    open(a.out, 'w').write('\n'.join(S) + '\n')
    print(f'\nالملف سليم ✓ — كتبنا {a.out}')


if __name__ == '__main__':
    main()
